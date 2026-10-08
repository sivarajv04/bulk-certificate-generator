"""Import recipient lists from CSV and Excel workbooks for bulk certificate jobs."""
import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.config import settings
from app.services.validation import RecipientValidationError, validate_recipient

SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}

_NAME_ALIASES = {
    "name", "studentname", "fullname", "recipientname", "participantname",
    "candidatename", "student", "participant", "candidate", "full_name",
    "student_name", "recipient_name", "participant_name", "candidate_name",
}
_EMAIL_ALIASES = {"email", "emailaddress", "mail", "studentemail", "student_email", "email_address"}


@dataclass(frozen=True)
class ImportedRow:
    row_number: int
    values: list[Any]
    name: str | None
    email: str | None
    valid: bool
    error: str | None


def _clean_header(value: Any, index: int) -> str:
    text = "" if value is None else str(value).strip()
    return text or f"Column {index + 1}"


def _normalized_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def _suggest_column(headers: list[str], aliases: set[str]) -> int | None:
    normalized = [_normalized_header(h) for h in headers]
    for i, value in enumerate(normalized):
        if value in {_normalized_header(a) for a in aliases}:
            return i
    for i, value in enumerate(normalized):
        if any(alias in value for alias in {_normalized_header(a) for a in aliases}):
            return i
    return None


def _read_csv(data: bytes) -> tuple[list[str], list[list[Any]]]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Could not decode the CSV file. Save it as UTF-8 and try again.")

    reader = csv.reader(io.StringIO(text))
    try:
        raw_headers = next(reader)
    except StopIteration:
        raise ValueError("The CSV file is empty") from None

    headers = [_clean_header(value, i) for i, value in enumerate(raw_headers)]
    rows = [list(row) + [None] * (len(headers) - len(row)) for row in reader]
    return headers, [row[: len(headers)] for row in rows]


def _read_excel(data: bytes) -> tuple[list[str], list[list[Any]]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if not workbook.sheetnames:
            raise ValueError("The Excel workbook has no worksheets")
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        try:
            raw_headers = list(next(iterator))
        except StopIteration:
            raise ValueError("The Excel worksheet is empty") from None
        headers = [_clean_header(value, i) for i, value in enumerate(raw_headers)]
        rows = []
        for raw_row in iterator:
            row = list(raw_row)
            row += [None] * (len(headers) - len(row))
            rows.append(row[: len(headers)])
        return headers, rows
    finally:
        workbook.close()


def parse_recipient_file(
    *,
    filename: str,
    data: bytes,
    name_column: int | None = None,
    email_column: int | None = None,
) -> tuple[list[str], int | None, int | None, list[ImportedRow]]:
    """Parse a CSV/XLSX file and validate each row without creating a job."""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported file type. Upload a .csv, .xlsx, or .xlsm file.")
    if len(data) > settings.max_import_file_bytes:
        limit_mb = settings.max_import_file_bytes / (1024 * 1024)
        raise ValueError(f"File is too large. Maximum size is {limit_mb:.1f} MB.")

    if suffix == ".csv":
        headers, raw_rows = _read_csv(data)
    else:
        try:
            headers, raw_rows = _read_excel(data)
        except Exception as exc:
            if isinstance(exc, ValueError):
                raise
            raise ValueError(f"Could not read the Excel file: {exc}") from exc

    if not headers or not any(h.strip() for h in headers):
        raise ValueError("The first row must contain column headers.")

    if len(raw_rows) > settings.max_recipients:
        raise ValueError(
            f"The file contains {len(raw_rows)} data rows; the maximum is {settings.max_recipients}."
        )

    if name_column is None:
        name_column = _suggest_column(headers, _NAME_ALIASES)
    if email_column == -1:
        email_column = None
    elif email_column is None:
        email_column = _suggest_column(headers, _EMAIL_ALIASES)

    if name_column is None:
        # Return the rows as a preview with no validation so the UI can show the
        # available columns and let the user choose the name column.
        return headers, None, email_column, [
            ImportedRow(
                row_number=offset,
                values=values,
                name=None,
                email=None,
                valid=False,
                error="Select the name column to validate this row",
            )
            for offset, values in enumerate(raw_rows, start=2)
            if not all(value is None or str(value).strip() == "" for value in values)
        ]
    if name_column < 0 or name_column >= len(headers):
        raise ValueError("Selected name column is out of range.")
    if email_column is not None and (email_column < 0 or email_column >= len(headers)):
        raise ValueError("Selected email column is out of range.")

    imported: list[ImportedRow] = []
    for offset, values in enumerate(raw_rows, start=2):
        def value_at(index: int | None) -> Any:
            return values[index] if index is not None and index < len(values) else None

        raw_name = value_at(name_column)
        raw_email = value_at(email_column)

        # Completely blank rows are ignored instead of becoming failed certificates.
        if all(value is None or str(value).strip() == "" for value in values):
            continue

        try:
            valid = validate_recipient(raw_name, raw_email, settings.max_name_length)
        except RecipientValidationError as exc:
            imported.append(
                ImportedRow(
                    row_number=offset,
                    values=values,
                    name=str(raw_name).strip() if raw_name is not None else None,
                    email=str(raw_email).strip() if raw_email is not None else None,
                    valid=False,
                    error=str(exc),
                )
            )
        else:
            imported.append(
                ImportedRow(
                    row_number=offset,
                    values=values,
                    name=valid.name,
                    email=valid.email,
                    valid=True,
                    error=None,
                )
            )

    return headers, name_column, email_column, imported
